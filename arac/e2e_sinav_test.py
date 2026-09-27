# Sınav Hazırla modülü uçtan uca tarayıcı testi.
#   ana sayfa -> yeni sınav sihirbazı -> düzenleyicide tüm soru türleri (görsel, biçim, yapıştırma, boşluk yap)
#   -> otomatik kayıt + sayfa yenileme -> geri al / yinele -> önizleme: A–D grupları ve cevap anahtarı doğrulaması
#   -> PDF (baskı düzeni) -> Word (.docx, python-docx ile) -> yedek dosyası dışa / içe aktarma
#   -> test sınavı -> optik okuyucuya aktarma -> optik ekranı anahtarları gösteriyor mu
# Kullanım:
#   npm run build && npx vite preview   ->   python3 arac/e2e_sinav_test.py [http://localhost:4173/] [çıktı klasörü]
import json, re, struct, sys, tempfile, time, zlib
from pathlib import Path
from playwright.sync_api import sync_playwright

URL = sys.argv[1] if len(sys.argv) > 1 else "http://localhost:4173/"
CIKTI = Path(sys.argv[2] if len(sys.argv) > 2 else tempfile.mkdtemp(prefix="sinav-e2e-"))
CIKTI.mkdir(parents=True, exist_ok=True)
HARF = "ABCDE"

hatalar = []
def kontrol(kosul, mesaj):
    print(("  ✔ " if kosul else "  ✘ ") + mesaj)
    if not kosul:
        hatalar.append(mesaj)


def png_yaz(yol, g=480, y=300):
    """Test görseli: yatay renk geçişli PNG (harici kütüphane gerektirmez)"""
    satirlar = b""
    for j in range(y):
        satirlar += b"\x00" + bytes(v for i in range(g) for v in ((i * 255) // g, (j * 255) // y, 160))
    def blok(tur, veri):
        return struct.pack(">I", len(veri)) + tur + veri + struct.pack(">I", zlib.crc32(tur + veri) & 0xFFFFFFFF)
    Path(yol).write_bytes(b"\x89PNG\r\n\x1a\n" + blok(b"IHDR", struct.pack(">IIBBBBB", g, y, 8, 2, 0, 0, 0))
                          + blok(b"IDAT", zlib.compress(satirlar, 9)) + blok(b"IEND", b""))


GORSEL = CIKTI / "deneme.png"
png_yaz(GORSEL)

# çoktan seçmeli sorular: (kök, şıklar, doğru şık indeksi)
MC = [
    ("Türkiye'nin başkenti neresidir?", ["İstanbul", "Ankara", "İzmir", "Bursa"], 1),
    ("12 × 12 işleminin sonucu kaçtır?", ["124", "132", "144", "154"], 2),
    ("Hangisi bir asal sayıdır?", ["21", "27", "33", "29"], 3),
    ("Suyun kaynama sıcaklığı deniz seviyesinde kaç °C'dir?", ["100", "90", "80", "120"], 0),
]
DY = [("Güneş bir yıldızdır.", True), ("Ay kendi ışığını üretir.", False), ("Dünya Güneş'in etrafında döner.", True)]
ES = [("Newton", "Yerçekimi"), ("Pasteur", "Aşı"), ("Edison", "Ampul")]

konsol = []
KAYIT = lambda s: s.wait_for_selector(".sh-kayit.kaydedildi", timeout=8000)


def duzenlenebilire_yaz(s, loc, metin, degistir=False):
    loc.click()
    if degistir:                       # hazır yönergenin yerine yaz
        s.keyboard.press("Control+a")
    s.keyboard.type(metin)


with sync_playwright() as p:
    tarayici = p.chromium.launch()
    ctx = tarayici.new_context(viewport={"width": 1440, "height": 900}, accept_downloads=True)
    s = ctx.new_page()
    s.on("console", lambda m: konsol.append(f"{m.type}: {m.text}") if m.type == "error" else None)
    s.on("pageerror", lambda e: konsol.append(f"pageerror: {e}"))

    print("1) Ana sayfa ve sihirbaz")
    s.goto(URL)
    s.wait_for_selector("a.modul-karti.m-sinav")
    kontrol(s.locator("a.modul-karti").count() == 2, "ana sayfada iki modül kartı (sınav, optik)")
    s.click("a.modul-karti.m-sinav")
    s.wait_for_selector(".sh-sablon")
    kontrol(s.url.endswith("#/sinav"), "Sınav hazırla -> #/sinav")
    s.click("button:has-text('Yeni sınav')")
    s.click(".pencere .sh-sablon:has-text('Boş sayfa')")
    s.click(".pencere [aria-label='Grup sayısı'] button:has-text('A–D')")
    s.fill("input[placeholder^='Örn. Atatürk']", "Cumhuriyet Anadolu Lisesi")
    s.fill("input[placeholder='Örn. Matematik']", "Fen Bilimleri")
    s.fill("input[placeholder='Örn. 9. Sınıf']", "9/B")
    s.fill("input[placeholder^='Örn. 1. Dönem']", "1. Dönem 2. Yazılı")
    s.click("button:has-text('Sınavı oluştur')")
    s.wait_for_selector(".sh-bos-kagit")
    sinav_url = s.url
    kontrol(re.search(r"#/sinav/[0-9a-z]+$", sinav_url) is not None, "düzenleyici açıldı: " + sinav_url.split("#")[1])

    print("2) Sorular")
    s.fill(".sh-toplu input", "4")
    s.click(".sh-toplu button:has-text('çoktan seçmeli')")
    s.wait_for_selector(".sh-oge-coktan")
    kontrol(s.locator(".sh-oge-coktan").count() == 4, "toplu ekleme: 4 çoktan seçmeli")
    for i, (kok, siklar, dogru) in enumerate(MC):
        o = s.locator(".sh-oge-coktan").nth(i)
        duzenlenebilire_yaz(s, o.locator(".sh-soru-metni"), kok)
        for j, t in enumerate(siklar):
            duzenlenebilire_yaz(s, o.locator(".sh-sik-metin").nth(j), t)
        o.locator(".sh-sik-harf").nth(dogru).click()
    kontrol(s.locator(".sh-oge-coktan .sh-sik.dogru").count() == 4, "doğru şıklar işaretlendi")

    # kalın biçim: ilk sorunun kökünde "başkenti" kelimesi
    ilk = s.locator(".sh-oge-coktan").nth(0).locator(".sh-soru-metni")
    ilk.click()
    s.evaluate("""() => { const el = document.activeElement; const t = el.firstChild; const i = t.textContent.indexOf('başkenti');
      const r = document.createRange(); r.setStart(t, i); r.setEnd(t, i + 8); const g = getSelection(); g.removeAllRanges(); g.addRange(r) }""")
    s.click(".sh-arac button[aria-label='Kalın']")
    kontrol("<b>başkenti</b>" in ilk.inner_html(), "araç çubuğu: kalın")

    def menuden_ekle(ad):
        s.click(".sh-arac button:has-text('Soru ekle')")
        s.click(f".sh-menu [role=menuitem]:has-text('{ad}')")
        s.wait_for_timeout(150)

    # doğru / yanlış
    menuden_ekle("Doğru / Yanlış")
    o = s.locator(".sh-oge-dy").last
    duzenlenebilire_yaz(s, o.locator(".sh-soru-metni"), "Aşağıdaki ifadelerin başına doğru ise D, yanlış ise Y yazınız.", degistir=True)
    while o.locator(".sh-madde").count() > len(DY):
        o.locator("button[title='İfadeyi sil']").last.click()
    for j, (t, d) in enumerate(DY):
        duzenlenebilire_yaz(s, o.locator(".sh-madde-metin").nth(j), t)
        o.locator(".sh-madde").nth(j).locator("button", has_text="D" if d else "Y").click()

    # boşluk doldurma: [köşeli parantez] ve "Boşluk yap" düğmesi
    menuden_ekle("Boşluk doldurma")
    o = s.locator(".sh-oge-bosluk").last
    duzenlenebilire_yaz(s, o.locator(".sh-soru-metni"), "Aşağıdaki cümlelerdeki boşlukları doldurunuz.", degistir=True)
    while o.locator(".sh-madde").count() > 2:
        o.locator("button[title='Cümleyi sil']").last.click()
    while o.locator(".sh-madde").count() < 2:
        o.locator("button:has-text('Cümle ekle')").click()
    duzenlenebilire_yaz(s, o.locator(".sh-madde-metin").nth(0), "Türkiye'nin başkenti [Ankara]'dır.")
    duzenlenebilire_yaz(s, o.locator(".sh-madde-metin").nth(1), "Suyun formülü H2O olarak yazılır.")
    s.evaluate("""() => { const el = document.activeElement; const t = el.firstChild; const i = t.textContent.indexOf('H2O');
      const r = document.createRange(); r.setStart(t, i); r.setEnd(t, i + 3); const g = getSelection(); g.removeAllRanges(); g.addRange(r) }""")
    s.click(".sh-arac button:has-text('Boşluk yap')")
    s.locator(".sh-kagit-baslik .sh-b-okul").click()   # odak çıkınca [..] boşluğa dönüşür
    s.wait_for_timeout(200)
    kontrol(o.locator("span.bosluk").count() == 2, "boşluk doldurma: [köşeli] + 'Boşluk yap' = 2 boşluk")
    o.locator(".sh-soru-no").click()
    kontrol(s.locator(".sh-ozellik label:has-text('Kelime havuzunu göster') input").is_checked(), "kelime havuzu varsayılan olarak açık")
    s.fill(".sh-ozellik input[placeholder^='Virgülle ayırın']", "İstanbul")

    # eşleştirme
    menuden_ekle("Eşleştirme")
    o = s.locator(".sh-oge-eslestirme").last
    duzenlenebilire_yaz(s, o.locator(".sh-soru-metni"), "Bilim insanlarını buluşlarıyla eşleştiriniz.", degistir=True)
    while o.locator(".sh-cift").count() > len(ES):
        o.locator("button[title='Çifti sil']").last.click()
    while o.locator(".sh-cift").count() < len(ES):
        o.locator("button:has-text('Çift ekle')").click()
    for j, (a, b) in enumerate(ES):
        duzenlenebilire_yaz(s, o.locator(".sh-cift").nth(j).locator(".sh-cift-metin").nth(0), a)
        duzenlenebilire_yaz(s, o.locator(".sh-cift").nth(j).locator(".sh-cift-metin").nth(1), b)

    # açık uçlu + görsel
    menuden_ekle("Açık uçlu")
    o = s.locator(".sh-oge-klasik").last
    duzenlenebilire_yaz(s, o.locator(".sh-soru-metni"), "Yukarıdaki grafikte görülen değişimi açıklayınız.")
    o.locator(".sh-oge-arac input[type=file]").set_input_files(str(GORSEL))
    o.locator(".sh-gorsel img").wait_for(timeout=8000)
    kontrol(o.locator(".sh-gorsel img").count() == 1, "görsel eklendi (açık uçlu soru)")
    o.locator(".sh-gorsel").hover()
    o.locator(".sh-gorsel-arac button[title='Metnin yanında']").click()
    kontrol(o.locator(".sh-gorsel.konum-yan").count() == 1, "görsel metnin yanına alındı")
    # şıkka görsel: 4. sorunun A şıkkı
    m4 = s.locator(".sh-oge-coktan").nth(3)
    m4.locator(".sh-soru-no").click()
    m4.locator(".sh-sik").nth(0).locator("input[type=file]").set_input_files(str(GORSEL))
    m4.locator(".sh-sik").nth(0).locator(".sh-gorsel img").wait_for(timeout=8000)
    kontrol(True, "şıkka görsel eklendi")

    # kısa cevaplı + yapıştırma temizliği
    menuden_ekle("Kısa cevaplı")
    o = s.locator(".sh-oge-klasik").last
    o.locator(".sh-soru-metni").click()
    s.evaluate("""() => { const dt = new DataTransfer();
      dt.setData('text/html', '<p style="color:red;font-size:40px" onclick="x()">Işığın <b>boşluktaki</b> hızı kaç km/s\\'dir?<script>window.__kotu=1<\\/script></p>');
      dt.setData('text/plain', 'Işığın boşluktaki hızı kaç km/s\\'dir?');
      document.activeElement.dispatchEvent(new ClipboardEvent('paste', { clipboardData: dt, bubbles: true, cancelable: true })) }""")
    html = o.locator(".sh-soru-metni").inner_html()
    kontrol("boşluktaki" in html and "script" not in html and "color" not in html and "onclick" not in html,
            "yapıştırılan HTML temizlendi (biçim korunur, script/stil atılır)")

    # bölüm başlığı en üste
    s.locator(".sh-ekle-cizgi .sh-ekle-arti").first.click(force=True)
    s.click(".sh-menu [role=menuitem]:has-text('Bölüm başlığı')")
    b = s.locator(".sh-oge-bolum").first
    b.locator(".sh-bolum-baslik").click()
    s.keyboard.type("A) Aşağıdaki soruları cevaplayınız.")
    kontrol(s.locator(".sh-kagit .sh-oge").first.get_attribute("class").find("sh-oge-bolum") >= 0, "bölüm başlığı en üste eklendi")

    s.locator(".sh-b-bilgi input").nth(0).fill("14.11.2026")
    s.locator(".sh-b-bilgi input").nth(1).fill("40 dakika")
    s.locator(".sh-b-bilgi input").nth(2).fill("Ayşe Yılmaz")
    s.locator(".sh-kagit-baslik").click(position={"x": 4, "y": 4})
    KAYIT(s)
    ust = s.locator(".sh-ust-alt").inner_text()
    kontrol("9 soru" in ust, f"üst bilgi: {ust.splitlines()[0]}")
    s.screenshot(path=str(CIKTI / "01-duzenleyici.png"))

    print("3) Geri al / yinele, kayıt kalıcılığı")
    n0 = s.locator(".sh-oge").count()
    s.locator(".sh-oge-klasik").last.locator(".sh-soru-no").click()
    s.locator(".sh-oge-klasik").last.locator("button[title='Sil']").click()
    kontrol(s.locator(".sh-oge").count() == n0 - 1, "soru silindi")
    s.keyboard.press("Control+z")
    kontrol(s.locator(".sh-oge").count() == n0, "Ctrl+Z geri getirdi")
    s.keyboard.press("Control+y")
    kontrol(s.locator(".sh-oge").count() == n0 - 1, "Ctrl+Y yeniden sildi")
    s.keyboard.press("Control+z")
    KAYIT(s)
    s.reload()
    s.wait_for_selector(".sh-oge")
    kontrol(s.locator(".sh-oge").count() == n0, "yenileme sonrası tüm öğeler duruyor")
    kontrol(s.locator(".sh-gorsel img").count() == 2, "yenileme sonrası görseller duruyor (IndexedDB)")
    kontrol("<b>başkenti</b>" in s.locator(".sh-oge-coktan").nth(0).locator(".sh-soru-metni").inner_html(), "yenileme sonrası biçim duruyor")

    print("4) Önizleme, gruplar ve cevap anahtarı")
    s.click("a:has-text('Önizle ve yazdır')")
    s.wait_for_selector(".sh-sayfalar .bs-sayfa:not(.bs-yukleniyor)")
    sekmeler = s.locator(".sh-sekmeler [role=tab]")
    kontrol(sekmeler.count() == 5, "sekmeler: A, B, C, D + cevap anahtarı")
    gruplar = []
    for g in range(4):
        sekmeler.nth(g).click()
        s.wait_for_timeout(250)
        s.wait_for_selector(".sh-sayfalar .bs-sayfa:not(.bs-yukleniyor)")
        veri = s.evaluate("""() => Array.from(document.querySelectorAll('.sh-sayfalar .bs-sayfa .bs-soru')).map(e => ({
          no: parseInt(e.querySelector('.bs-no').textContent),
          kok: e.querySelector('.bs-metin-ic')?.textContent.trim(),
          siklar: Array.from(e.querySelectorAll('.bs-sik .bs-sik-metin')).map(x => x.textContent.trim()),
          sol: Array.from(e.querySelectorAll('.bs-es-sol')).map(x => x.textContent.trim()).filter(Boolean),
          sag: Array.from(e.querySelectorAll('.bs-es-sag')).map(x => x.textContent.trim()).filter(Boolean),
          grupHarf: e.closest('.bs-sayfa').dataset.grup,
        }))""")
        gruplar.append(veri)
        s.screenshot(path=str(CIKTI / f"02-onizleme-{HARF[g]}.png"))
    kontrol(all(len(v) == 9 for v in gruplar), "her grupta 9 soru")
    kontrol(all(all(q["grupHarf"] == HARF[g] for q in v) for g, v in enumerate(gruplar)), "sayfalar doğru grup harfiyle işaretli")
    siralar = [[q["kok"] for q in v] for v in gruplar]
    kontrol(len({tuple(x) for x in siralar}) == 4, "4 grubun soru sıraları birbirinden farklı")
    sekmeler.nth(4).click()
    s.wait_for_selector(".bs-anahtar-grup")
    s.screenshot(path=str(CIKTI / "03-cevap-anahtari.png"))
    anahtar = s.evaluate("""() => Array.from(document.querySelectorAll('.bs-anahtar-grup')).map(g => ({
      mc: Object.fromEntries(Array.from(g.querySelectorAll('.bs-anahtar-izgara span')).map(x => [x.querySelector('i').textContent, x.querySelector('b').textContent])),
      diger: Object.fromEntries(Array.from(g.querySelectorAll('.bs-anahtar-tablo tbody tr')).map(tr => [tr.querySelector('.no').textContent, tr.children[1].textContent])),
    }))""")
    dogru_metin = {k: siklar[d] for k, siklar, d in MC}
    es_dogru = dict(ES)
    tutarli = True
    for g, v in enumerate(gruplar):
        for q in v:
            if q["siklar"]:
                harf = anahtar[g]["mc"].get(str(q["no"]))
                if not harf or q["siklar"][HARF.index(harf)] != dogru_metin[q["kok"]]:
                    tutarli = False; print("    anahtar hatası", HARF[g], q)
            elif q["sol"]:
                cift = anahtar[g]["diger"].get(str(q["no"]), "")
                for m in re.finditer(r"(\d+)-([a-z])", cift):
                    sol = q["sol"][int(m.group(1)) - 1]
                    sag = q["sag"]["abcdefghij".index(m.group(2))]
                    if es_dogru[sol] != sag:
                        tutarli = False; print("    eşleştirme hatası", HARF[g], sol, sag)
            elif "ifadelerin" in (q["kok"] or ""):
                kontrol(anahtar[g]["diger"].get(str(q["no"]), "").replace(" ", "") == "a)Db)Yc)D", f"{HARF[g]} grubu D/Y anahtarı") if g == 0 else None
            elif "boşlukları" in (q["kok"] or ""):
                t = anahtar[g]["diger"].get(str(q["no"]), "")
                if not ("Ankara" in t and "H2O" in t):
                    tutarli = False; print("    boşluk anahtarı hatası", t)
    kontrol(tutarli, "4 grubun cevap anahtarı kâğıttaki şık/eşleştirme sırasıyla birebir tutarlı")
    kontrol(s.locator(".bs-anahtar-sayfa .bs-optik-satir").count() == 0, "karma sınavda 'optik anahtar' satırı yok")

    print("5) PDF (yazdırma düzeni)")
    sekmeler.nth(0).click()
    s.evaluate("() => { window.__yazdir = 0; window.print = () => { window.__yazdir++ } }")
    s.locator(".sh-bolunmus-ok").click()
    s.click(".sh-menu [role=menuitem]:has-text('Gruplar ve cevap anahtarı')")
    s.wait_for_function("() => window.__yazdir > 0", timeout=20000)
    pdf = CIKTI / "sinav.pdf"
    s.pdf(path=str(pdf), prefer_css_page_size=True)
    s.wait_for_function("() => !document.documentElement.classList.contains('sh-yazdir')", timeout=5000)
    from pypdf import PdfReader
    r = PdfReader(str(pdf))
    boyut = r.pages[0].mediabox
    kontrol(abs(float(boyut.width) - 595.3) < 2 and abs(float(boyut.height) - 841.9) < 2, f"PDF sayfası A4 ({float(boyut.width):.0f}×{float(boyut.height):.0f} pt)")
    metinler = [pg.extract_text() or "" for pg in r.pages]
    kontrol(len(r.pages) >= 5, f"PDF: {len(r.pages)} sayfa (4 grup + anahtar)")
    grup_sayfalari = [next((i for i, t in enumerate(metinler) if "GRUBU" in t and h in t), -1) for h in HARF[:4]]
    kontrol(all(i >= 0 for i in grup_sayfalari), "PDF'te A, B, C, D grupları var")
    ilk_anahtar = next((i for i, t in enumerate(metinler) if "CEVAP ANAHTARI" in t), -1)
    kontrol(ilk_anahtar > max(grup_sayfalari), f"PDF: cevap anahtarı grupların ardından ({ilk_anahtar + 1}. sayfadan itibaren)")
    kontrol(all("GRUBU" in metinler[i] or "puan" in metinler[i] for i in range(len(metinler))), "PDF: boş sayfa yok")
    kontrol(not any("Önizleme" in t or "Sınav ayarları" in t.upper() for t in metinler), "PDF'e arayüz öğeleri karışmadı")
    import subprocess
    subprocess.run(["pdftoppm", "-r", "60", "-png", "-f", "1", "-l", "1", str(pdf), str(CIKTI / "04-pdf")], check=False)

    print("6) Word (.docx)")
    with s.expect_download(timeout=30000) as d:
        s.click(".sh-ust-dugme:has-text('Word')")
    docx_yol = CIKTI / d.value.suggested_filename
    d.value.save_as(str(docx_yol))
    kontrol(docx_yol.suffix == ".docx" and docx_yol.stat().st_size > 5000, f"Word indirildi: {docx_yol.name} ({docx_yol.stat().st_size // 1024} KB)")
    import docx as pydocx
    w = pydocx.Document(str(docx_yol))
    tum = "\n".join(p.text for p in w.paragraphs) + "\n" + "\n".join(c.text for t in w.tables for row in t.rows for c in row.cells)
    kontrol(all(k in tum for k, _, _ in MC), "Word: tüm çoktan seçmeli kökleri var")
    kontrol("Newton" in tum and "Yerçekimi" in tum, "Word: eşleştirme tablosu var")
    kontrol("CEVAP ANAHTARI" in tum, "Word: cevap anahtarı bölümü var")
    kontrol("Ankara" not in "\n".join(p.text for p in w.paragraphs[:200] if "başkenti" in p.text and "____" in p.text) and
            any("____" in p.text for p in w.paragraphs), "Word: öğrenci kâğıdında boşluklar çizgi, cevap yazmıyor")
    xml = w.element.xml
    kontrol(xml.count("<wp:inline") == 4 and xml.count("<wp:anchor") == 4, f"Word: her grupta şık görseli (satır içi) + metin yanı görsel (kayan): {xml.count('<wp:inline')} + {xml.count('<wp:anchor')}")
    kontrol("İstanbul" in tum and all(k in tum for k in ("Ankara", "H2O")), "Word: kelime havuzu (çeldirici dahil)")
    kontrol(len(w.sections) == 5, f"Word: {len(w.sections)} bölüm (4 grup + anahtar)")
    kontrol(abs(w.sections[0].page_width.mm - 210) < 1 and abs(w.sections[0].page_height.mm - 297) < 1, "Word: A4 sayfa")
    kontrol(any(r.bold for p in w.paragraphs for r in p.runs if r.text == "başkenti"), "Word: kalın biçim korundu")

    print("7) İki sütun, grup sayısı, sınavlarım")
    s.click(".sh-oniz-yan [aria-label='Sütun'] button:has-text('İki')")
    s.wait_for_timeout(400)
    s.wait_for_selector(".sh-sayfalar .bs-sayfa:not(.bs-yukleniyor)")
    dolu = s.evaluate("() => Array.from(document.querySelectorAll('.sh-sayfalar .bs-sayfa')[0].querySelectorAll('.bs-sutun')).map(c => c.children.length)")
    kontrol(len(dolu) == 2 and all(n > 0 for n in dolu), f"iki sütun: ilk sayfada sütunlar dolu {dolu}")
    s.screenshot(path=str(CIKTI / "02-onizleme-iki-sutun.png"))
    with s.expect_download(timeout=30000) as d:
        s.click(".sh-ust-dugme:has-text('Word')")
    d.value.save_as(str(CIKTI / "iki-sutun.docx"))
    w2 = pydocx.Document(str(CIKTI / "iki-sutun.docx"))
    kontrol('w:num="2"' in w2.element.xml and len(w2.sections) == 9, f"Word iki sütun: {len(w2.sections)} bölüm, sütunlu gövde")
    s.click(".sh-oniz-yan [aria-label='Sütun'] button:has-text('Tek')")
    s.click(".sh-oniz-yan [aria-label='Grup sayısı'] button:has-text('Tek')")
    s.wait_for_timeout(300)
    kontrol(s.locator(".sh-sekmeler [role=tab]").count() == 2 and "Sınav kâğıdı" in s.locator(".sh-sekmeler").inner_text(), "tek grup: 'Sınav kâğıdı' + anahtar")
    s.click(".sh-oniz-yan [aria-label='Grup sayısı'] button:has-text('A–D')")
    s.goto(URL + "#/sinav")
    s.wait_for_selector(".sh-sinav-karti:not(.yeni)")
    kontrol(s.locator(".sh-sinav-karti:not(.yeni)").count() == 1, "sınavlarım: 1 sınav")
    s.locator(".sh-sinav-karti:not(.yeni)").first.hover()
    s.locator(".sh-sinav-karti:not(.yeni) button[aria-label='Diğer işlemler']").first.click()
    with s.expect_download() as d:
        s.click(".sh-menu [role=menuitem]:has-text('Yedek')")
    yedek = CIKTI / d.value.suggested_filename
    d.value.save_as(str(yedek))
    y = json.loads(yedek.read_text())
    kontrol(yedek.suffix == ".sinav" and len(y.get("gorseller", {})) == 2, "yedek dosyası (.sinav) görselleriyle birlikte indirildi")
    s.locator("input[type=file][accept*='.sinav']").set_input_files(str(yedek))
    s.wait_for_function("() => document.querySelectorAll('.sh-sinav-karti:not(.yeni)').length === 2", timeout=8000)
    kontrol(True, "yedek içe aktarıldı (2 sınav)")
    s.screenshot(path=str(CIKTI / "05-sinavlarim.png"))

    print("8) Test sınavı -> optik okuyucu")
    s.click("button:has-text('Yeni sınav')")
    s.click(".pencere .sh-sablon:has-text('Test')")
    s.fill(".pencere input[type=number]", "12")
    s.click(".pencere [aria-label='Şık sayısı'] button:has-text('A–D')")
    s.click(".pencere [aria-label='Grup sayısı'] button:has-text('A–B')")
    s.fill("input[placeholder^='Örn. 1. Dönem']", "Deneme Testi")
    s.click("button:has-text('Sınavı oluştur')")
    s.wait_for_selector(".sh-oge-coktan")
    n = s.locator(".sh-oge-coktan").count()
    kontrol(n == 12, f"test şablonu: {n} soru")
    for i in range(n):
        o = s.locator(".sh-oge-coktan").nth(i)
        for j in range(4):
            o.locator(".sh-sik-metin").nth(j).click()
            s.keyboard.type(f"S{i + 1}-{'ABCD'[j]}")
        o.locator(".sh-sik-harf").nth(i % 4).click()
    KAYIT(s)
    s.click("a:has-text('Önizle ve yazdır')")
    s.wait_for_selector(".sh-optik-kart button:has-text('Optik okuyucuya aktar')")
    s.locator(".sh-sekmeler [role=tab]").last.click()
    optik_satirlari = s.locator(".bs-optik-satir code").all_inner_texts()
    kontrol(len(optik_satirlari) == 2 and optik_satirlari[0] == "ABCD" * 3, f"cevap anahtarı optik satırları: {optik_satirlari}")
    s.click(".sh-optik-kart button:has-text('Optik okuyucuya aktar')")
    s.click(".pencere button:has-text('Aktar ve okumaya başla')")
    s.wait_for_function("() => location.hash === '#/optik'")
    s.wait_for_timeout(600)
    durum = s.evaluate("() => JSON.parse(localStorage.getItem('optik-okuyucu.sinav.v1'))")
    ak = {k: "".join("ABCDE"[x] for x in v[:12]) for k, v in durum["anahtarlar"].items()}
    kontrol(durum["ayar"]["soruSayisi"] == 12 and ak.get("A") == optik_satirlari[0] and ak.get("B") == optik_satirlari[1],
            f"optik kaydı: 12 soru, anahtarlar {ak}")
    kontrol(durum["ekran"] == "okut" and s.locator(".logo:has-text('Optik Okuyucu')").count() == 1, "optik okuyucu 'okut' ekranında açıldı")
    s.screenshot(path=str(CIKTI / "06-optik.png"))
    s.click("a.marka")
    s.wait_for_selector("a.modul-karti.m-optik")
    kontrol("Deneme Testi" in s.locator("a.modul-karti.m-optik").inner_text(), "ana sayfa: optik kartında aktarılan sınav (okumaya hazır)")
    kontrol(s.locator(".ana-son-sinavlar a, .son-sinav").count() >= 1 or "Deneme Testi" in s.locator("body").inner_text(), "ana sayfa: son sınavlar")
    s.screenshot(path=str(CIKTI / "07-ana-sayfa.png"))

    tarayici.close()

kontrol(not konsol, "tarayıcı konsolunda hata yok" + ("" if not konsol else ": " + " | ".join(konsol[:5])))
print(f"\nÇıktılar: {CIKTI}")
print("SONUÇ:", "BAŞARILI" if not hatalar else f"{len(hatalar)} HATA")
sys.exit(1 if hatalar else 0)
