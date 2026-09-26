# Uçtan uca tarayıcı testi: ayarlar -> anahtar (fotoğraftan) -> öğrenciler -> kontrol pencereleri -> yenileme -> tekrar -> Excel -> e-posta
import base64, io, json, re, sys, time
from pathlib import Path
from playwright.sync_api import sync_playwright
import openpyxl

KOK = Path(sys.argv[1])
URL = sys.argv[2] if len(sys.argv) > 2 else "http://localhost:4173/"
B = json.loads((KOK / "beklenen.json").read_text())
N = B["N"]
SIK = "ABCDE"


def baslik(s):
    return " ".join(k[:1].upper() + k[1:] for k in s.replace("I", "ı").replace("İ", "i").lower().split()).replace("ı", "ı")


def tr_baslik(s):
    # Türkçe küçük/büyük harf dönüşümü
    kucuk = s.replace("I", "ı").replace("İ", "i").lower()
    out = []
    for k in kucuk.split():
        ilk = k[0]
        ilk = "İ" if ilk == "i" else "I" if ilk == "ı" else ilk.upper()
        out.append(ilk + k[1:])
    return " ".join(out)


hatalar = []
def kontrol(kosul, mesaj):
    print(("  ✔ " if kosul else "  ✘ ") + mesaj)
    if not kosul:
        hatalar.append(mesaj)


with sync_playwright() as p:
    tarayici = p.chromium.launch(args=["--use-fake-ui-for-media-stream", "--use-fake-device-for-media-stream"])
    ctx = tarayici.new_context(viewport={"width": 430, "height": 920}, accept_downloads=True, permissions=["camera"])
    sayfa = ctx.new_page()
    konsol = []
    sayfa.on("console", lambda m: konsol.append(f"{m.type}: {m.text}") if m.type in ("error",) else None)
    sayfa.on("pageerror", lambda e: konsol.append(f"pageerror: {e}"))
    sayfa.on("dialog", lambda d: d.accept())
    sayfa.goto(URL)

    print("1) Ayarlar")
    sayfa.fill('input[placeholder^="Örn."]', "E2E Test Sınavı")
    sayfa.fill('label:has-text("Soru sayısı") input', str(N))
    kontrol(sayfa.locator("text=Toplam:").inner_text().find("100") >= 0, "otomatik puan toplamı 100")
    sayfa.click("text=Devam: Cevap anahtarı")

    print("2) Cevap anahtarları")
    t0 = time.time()
    sayfa.wait_for_selector('button:has-text("Anahtarı kamerayla okut"):not([disabled])', timeout=180000)
    print(f"  okuyucu {time.time() - t0:.1f} sn'de hazır")
    for kit in ("A", "B"):
        sayfa.set_input_files('input[type=file]', str(KOK / "foto" / f"anahtar_{kit}.jpg"))
        sayfa.wait_for_selector("text=Anahtarı kaydet", timeout=60000)
        secili = sayfa.locator(".anahtar-satir").evaluate_all(
            "els => els.map(e => { const s = e.querySelector('.sik.secili'); return s ? s.textContent : null })")
        beklenen = [SIK[B["anahtarlar"][kit][str(q)]] for q in range(N)]
        kontrol(secili == beklenen, f"{kit} anahtarı 40 sorunun tamamı doğru okundu")
        kontrol(sayfa.locator(".pencere .sec.secili").first.inner_text() == kit, f"{kit} kitapçığı otomatik seçildi")
        kontrol(not sayfa.is_visible("text=Boş sorular"), "anahtarda eksik soru yok")
        sayfa.click("text=Anahtarı kaydet")
        sayfa.wait_for_selector(f"text=Kitapçık {kit}")
    sayfa.click("text=Devam: Öğrenci kâğıtlarını okut")

    print("3) Öğrenci kâğıtları (fotoğraftan)")
    dosyalar = [str(KOK / "foto" / o["dosya"]) for o in B["ogrenciler"]]
    sayfa.set_input_files('input[type=file]', dosyalar)
    pencere_sayisi = 0
    son = time.time()
    while time.time() - son < 120:
        if sayfa.is_visible("text=Onayla ve kaydet"):
            pencere_sayisi += 1
            # hangi öğrenci? numaradan bul
            no_input = sayfa.locator('.form-izgara label:has-text("Numara") input')
            ad_in = sayfa.locator('.form-izgara label:has-text("Ad") input').first
            okunan_no = no_input.input_value()
            okunan_ad = ad_in.input_value()
            ogr = next((o for o in B["ogrenciler"] if o["no"] == okunan_no), None) or \
                  next((o for o in B["ogrenciler"] if o["ad"] == okunan_ad), None)
            if ogr is None:
                # adı/numarası soru işaretli olabilir: kalıba göre eşleştir
                desen = re.compile("^" + re.escape(okunan_no).replace("\\?", ".") + "$")
                ogr = next((o for o in B["ogrenciler"] if desen.match(o["no"])), None)
            kontrol(ogr is not None, f"kontrol penceresindeki öğrenci tanındı ({okunan_ad} / {okunan_no})")
            if ogr is None:
                sayfa.click("text=Bu kâğıdı atla"); continue
            sorunlar = sayfa.locator(".sorun-listesi").inner_text()
            print(f"  · kontrol penceresi: {ogr['dosya']} -> {sorunlar.strip().splitlines()[:3]}")
            ad_in.fill(ogr["ad"])
            sayfa.locator('.form-izgara label:has-text("Soyad") input').fill(ogr["soyad"])
            no_input.fill(ogr["no"])
            for blok in sayfa.locator(".soru-kontrol").all():
                q = int(blok.locator(".soru-baslik").inner_text().split(".")[0]) - 1
                c = ogr["cevaplar"][str(q)]
                if c["tur"] == "cift":
                    blok.locator("button:has-text('Çift')").click()
                elif c["tur"] == "bos":
                    blok.locator("button:has-text('Boş')").click()
                else:
                    blok.locator(f"button.sec:text-is('{SIK[c['k']]}')").click()
            sayfa.click("text=Onayla ve kaydet")
            time.sleep(0.2)
            son = time.time()
            continue
        sayac = sayfa.locator(".sayac").inner_text()
        if sayac.startswith(f"{len(dosyalar)} ") and not sayfa.is_visible("text=Onayla ve kaydet"):
            break
        time.sleep(0.3)
    # Okunamayan (ör. çok uzaktan çekilmiş) kâğıtlar: uygulama uyarmalı; öğretmen yeniden çekip okutur
    okunan_sayi = int(sayfa.locator(".sayac").inner_text().split()[0])
    if okunan_sayi < len(dosyalar):
        # numara benzersiz: eksikleri numaradan bul (aynı adlı iki öğrenci olabilir)
        gorulen_nolar = set(sayfa.locator(".sonuc-listesi li .isim small").evaluate_all(
            "els => els.map(e => (e.textContent.match(/No (\\d+)/) || [])[1])"))
        eksikler = [o for o in B["ogrenciler"] if o["no"] not in gorulen_nolar]
        print(f"  · {len(eksikler)} kâğıt okunamadı (beklenen: düşük çözünürlüklü foto) -> yeniden çekiliyor: {[o['dosya'] for o in eksikler]}")
        import cv2, numpy as np
        yeni = []
        for o in eksikler:
            duz = cv2.imread(str(KOK / f"duz_{o['dosya'][:-4]}.png"))
            h, w = duz.shape[:2]
            tuval = np.full((int(h * 1.12), int(w * 1.12), 3), (90, 110, 140), np.uint8)
            M = cv2.getRotationMatrix2D((w / 2, h / 2), 3, 1.0); M[:, 2] += [w * 0.06, h * 0.06]
            foto = cv2.warpAffine(duz, M, (tuval.shape[1], tuval.shape[0]), dst=tuval, borderMode=cv2.BORDER_TRANSPARENT)
            yol = KOK / "foto" / f"yeniden_{o['dosya']}"
            cv2.imwrite(str(yol), foto, [cv2.IMWRITE_JPEG_QUALITY, 88])
            yeni.append(str(yol))
        sayfa.set_input_files('input[type=file]', yeni)
        t1 = time.time()
        while time.time() - t1 < 60 and not sayfa.locator(".sayac").inner_text().startswith(f"{len(dosyalar)} "):
            time.sleep(0.3)
    kontrol(sayfa.locator(".sayac").inner_text().startswith(f"{len(dosyalar)} "), f"{len(dosyalar)} öğrencinin hepsi listede")
    kontrol(pencere_sayisi >= 2, f"zor kâğıtlarda kontrol penceresi açıldı ({pencere_sayisi})")

    def liste_oku():
        return sayfa.locator(".sonuc-listesi li").evaluate_all(
            "els => els.map(e => [e.querySelector('.isim').childNodes[0].textContent.trim(), e.querySelector('.puan').textContent.trim(), e.querySelector('.isim small').textContent])")

    liste = liste_oku()
    beklenen = {(tr_baslik(o["ad"] + " " + o["soyad"]), o["no"], f"{o['puan']:.2f}".replace(".", ",")) for o in B["ogrenciler"]}
    gorulen = {(a, (re.search(r"No (\d+)", k) or [None, None])[1], p) for a, p, k in liste}
    kontrol(gorulen == beklenen, "tüm isimler, numaralar ve puanlar birebir doğru")
    if gorulen != beklenen:
        print("   eksik:", beklenen - gorulen)
        print("   fazla:", gorulen - beklenen)

    print("4) Sayfa yenileme sonrası kalıcılık")
    sayfa.reload()
    sayfa.wait_for_selector(".sonuc-listesi li")
    kontrol(len(liste_oku()) == len(dosyalar), "yenilemeden sonra liste duruyor")

    print("5) Aynı kâğıdı tekrar okutma")
    sayfa.wait_for_selector('button:has-text("Fotoğraftan okut"):not([disabled])', timeout=180000)
    sayfa.set_input_files('input[type=file]', str(KOK / "foto" / B["ogrenciler"][2]["dosya"]))
    sayfa.wait_for_selector("text=zaten okunmuş", timeout=60000)
    kontrol(len(liste_oku()) == len(dosyalar), "tekrar okutulan kâğıt listeye eklenmedi")

    print("6) Excel")
    sayfa.click("text=Bitti → Excel")
    with sayfa.expect_download() as indir:
        sayfa.click("text=Excel'i indir")
    yol = KOK / "indirilen.xlsx"
    indir.value.save_as(str(yol))
    wb = openpyxl.load_workbook(yol)
    kontrol(wb.sheetnames == ["Sonuçlar", "Sıralama", "Cevaplar", "Soru Analizi", "Cevap Anahtarı", "Bilgi"], f"Excel sayfaları: {wb.sheetnames}")
    ws = wb["Sonuçlar"]
    satirlar = list(ws.iter_rows(min_row=2, values_only=True))
    kontrol(len(satirlar) == len(dosyalar), "Excel'de öğrenci sayısı doğru")
    excel = {(str(r[1]), f"{r[2]} {r[3]}", r[4], r[5], r[6], r[7], round(r[9], 2)) for r in satirlar}
    bek = {(o["no"], tr_baslik(o["ad"] + " " + o["soyad"]), o["kitapcik"], o["d"], o["y"], o["b"], o["puan"]) for o in B["ogrenciler"]}
    kontrol(excel == bek, "Excel satırları (no, ad soyad, kitapçık, D/Y/B, puan) birebir doğru")
    if excel != bek:
        print("   eksik:", bek - excel); print("   fazla:", excel - bek)
    wc = wb["Cevaplar"]
    kontrol(wc.max_column == 4 + N, "Cevaplar sayfasında 40 soru sütunu var")
    wa = wb["Cevap Anahtarı"]
    anahtar_satir = {r[0]: "".join(r[1:]) for r in wa.iter_rows(min_row=2, values_only=True)}
    kontrol(anahtar_satir == {k: "".join(SIK[v[str(q)]] for q in range(N)) for k, v in B["anahtarlar"].items()}, "Excel'deki cevap anahtarları doğru")

    print("7) E-posta")
    yakalanan = {}
    def yakala(route):
        yakalanan["govde"] = route.request.post_data_json
        route.fulfill(status=200, content_type="application/json", body='{"tamam":true}')
    sayfa.route("**/api/eposta", yakala)
    sayfa.fill('input[type=email]', "ogretmen@okul.k12.tr")
    sayfa.click("text=📧 Gönder")
    sayfa.wait_for_selector("text=adresine gönderildi", timeout=30000)
    g = yakalanan.get("govde") or {}
    kontrol(g.get("kime") == "ogretmen@okul.k12.tr", "e-posta adresi sunucuya iletildi")
    ek = base64.b64decode(g.get("veri", ""))
    kontrol(ek[:2] == b"PK" and len(openpyxl.load_workbook(io.BytesIO(ek))["Sonuçlar"]["A"]) == len(dosyalar) + 1, "e-postadaki Excel eki geçerli ve tam")

    sayfa.screenshot(path=str(KOK / "sonuc_ekrani.png"), full_page=True)
    kontrol(not konsol, f"tarayıcı konsolunda hata yok {konsol[:3]}")
    tarayici.close()

print("\nSONUÇ:", "TÜM KONTROLLER GEÇTİ" if not hatalar else f"{len(hatalar)} HATA: {hatalar}")
