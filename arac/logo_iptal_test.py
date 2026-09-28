# Okul logosu ve optikte soru iptali: uçtan uca (tarayıcıda) doğrulama.
#
#  1) Düzenleyicide sol ve sağ logo eklenir -> önizleme, yazdırma (PDF) ve Word başlığında görünür; yeni sınavlara da gelir.
#  2) Sınav optik okuyucuya aktarılır, öğrenciler eklenir -> Sonuç ekranında soru iptali:
#     "Herkese doğru say" ve "Soruyu çıkar"; diğer kitapçıktaki karşılığı otomatik bulunur; puanlar ve Excel
#     (bağımsız hesapla karşılaştırılarak) birebir; geri alınınca eski puanlar.
#
# Kullanım: npm run build && npx vite preview  ->  python3 arac/logo_iptal_test.py [adres] [çıktı klasörü]
import json, sys, tempfile, zipfile, re, subprocess
from pathlib import Path
import numpy as np
import cv2
from pypdf import PdfReader
from playwright.sync_api import sync_playwright

sys.path.insert(0, str(Path(__file__).resolve().parent))
from optik_baski_test import sinav_uret

URL = next((a for a in sys.argv[1:] if a.startswith("http")), "http://localhost:4173/").split("#")[0]
CIKTI = Path(next((a for a in sys.argv[1:] if not a.startswith("http") and not a.startswith("--")), "") or tempfile.gettempdir() + "/logo-iptal")

hatalar = []
def kontrol(kosul, mesaj):
    print(("  ✔ " if kosul else "  ✘ ") + mesaj, flush=True)
    if not kosul:
        hatalar.append(mesaj)


def logo_yap(yol, genis=False):
    w, h = (600, 220) if genis else (400, 400)
    img = np.full((h, w, 3), 255, np.uint8)
    if genis:
        cv2.rectangle(img, (10, 10), (w - 10, h - 10), (140, 60, 20), 12)
        cv2.putText(img, "UNIVERSITE", (40, 135), cv2.FONT_HERSHEY_DUPLEX, 2.2, (140, 60, 20), 5)
    else:
        cv2.circle(img, (200, 200), 185, (30, 30, 200), 14)
        cv2.circle(img, (200, 200), 120, (30, 30, 200), -1)
        cv2.putText(img, "OKUL", (95, 225), cv2.FONT_HERSHEY_DUPLEX, 2.4, (255, 255, 255), 6)
    cv2.imwrite(str(yol), img)


def puan_hesapla(ogr, anahtar, N, iptal):
    """Bağımsız hesap: iptal = {soru: 'dogru' | 'cikar'} (öğrencinin kitapçığındaki numarayla)"""
    kalan = N - sum(1 for t in iptal.values() if t == "cikar")
    d = 0
    for q in range(N):
        if iptal.get(q) == "cikar":
            continue
        if iptal.get(q) == "dogru":
            d += 1
            continue
        c = ogr["cevaplar"][q]
        if c.get("t") == "c" and c["k"] == anahtar[q]:
            d += 1
    return round(d * 100 / kalan, 2)


def main():
    CIKTI.mkdir(parents=True, exist_ok=True)
    konsol = []
    logo1, logo2 = CIKTI / "logo_okul.png", CIKTI / "logo_uni.png"
    logo_yap(logo1)
    logo_yap(logo2, genis=True)
    N, K, G = 20, 5, 2
    yedek, ad, dogru = sinav_uret(N, K, G, 4242)
    (CIKTI / "sinav.sinav").write_text(json.dumps(yedek, ensure_ascii=False))

    with sync_playwright() as p:
        tarayici = p.chromium.launch()
        ctx = tarayici.new_context(viewport={"width": 1440, "height": 900}, accept_downloads=True)
        s = ctx.new_page()
        s.on("console", lambda m: konsol.append(f"{m.type}: {m.text}") if m.type == "error" else None)
        s.on("pageerror", lambda e: konsol.append(f"pageerror: {e}"))

        print("1) Okul logosu")
        s.goto(URL + "#/sinav")
        s.wait_for_selector(".sh-liste-sayfa")
        s.evaluate("() => { try { const k = 'optik-okuyucu.ogretmen'; const p = JSON.parse(localStorage.getItem(k) || '{}'); delete p.logoSol; delete p.logoSag; localStorage.setItem(k, JSON.stringify(p)) } catch {} }")
        s.locator("input[type=file][accept*='.sinav']").set_input_files(str(CIKTI / "sinav.sinav"))
        kart = s.locator(".sh-sinav-karti", has_text=ad)
        kart.first.wait_for(timeout=10000)
        sid = kart.first.locator("a.sh-kart-baglanti").get_attribute("href").split("#/sinav/")[1]
        s.goto(URL + f"#/sinav/{sid}")
        s.wait_for_selector(".sh-logo-yeri")
        kontrol(s.locator(".sh-logo-yeri").count() == 2, "başlıkta iki logo yeri (sol, sağ)")
        s.screenshot(path=str(CIKTI / "01-logo-bos.png"), clip={"x": 300, "y": 60, "width": 900, "height": 330})
        s.locator(".sh-logo-yeri").nth(0).locator("input[type=file]").set_input_files(str(logo1))
        s.wait_for_selector(".sh-logo-yeri.dolu img", timeout=10000)
        s.locator(".sh-logo-yeri").nth(1).locator("input[type=file]").set_input_files(str(logo2))
        s.wait_for_function("() => document.querySelectorAll('.sh-logo-yeri.dolu img').length === 2", timeout=10000)
        kontrol(True, "sol ve sağ logo eklendi")
        s.wait_for_timeout(900)
        s.screenshot(path=str(CIKTI / "02-logo-duzenleyici.png"), clip={"x": 300, "y": 60, "width": 900, "height": 330})
        profil = s.evaluate("() => JSON.parse(localStorage.getItem('optik-okuyucu.ogretmen') || '{}')")
        kontrol(bool(profil.get("logoSol")) and bool(profil.get("logoSag")), "logolar sonraki sınavlar için hatırlandı")

        s.goto(URL + f"#/sinav/{sid}/yazdir")
        s.wait_for_selector(".sh-sayfalar .bs-sayfa:not(.bs-yukleniyor)")
        s.wait_for_function("() => document.querySelectorAll('.sh-sayfalar .bs-sayfa')[0].querySelectorAll('.bs-logo img').length === 2", timeout=10000)
        kontrol(True, "önizlemede iki logo")
        s.locator(".sh-sayfalar .bs-sayfa").first.screenshot(path=str(CIKTI / "03-onizleme-sayfa.png"))

        # yazdır (her gruptan bir kâğıt) -> PDF'te logolar
        s.click(".sh-yazdir-dugme")
        s.wait_for_selector(".sh-yazdir-pencere")
        s.locator(".sh-yazdir-pencere .sh-secici button:has-text('Her gruptan 1')").click()
        s.wait_for_selector(".sh-yazdir-onay:not([disabled])", timeout=30000)
        s.evaluate("() => { window.__yazdir = 0; window.print = () => { window.__yazdir++ } }")
        s.click(".sh-yazdir-onay")
        s.wait_for_function("() => window.__yazdir > 0", timeout=60000)
        s.wait_for_function("() => Array.from(document.querySelectorAll('.bs-baski-alani .bs-logo img')).every(i => i.complete && i.naturalWidth > 0)", timeout=10000)
        pdf = CIKTI / "baski.pdf"
        s.pdf(path=str(pdf), prefer_css_page_size=True)
        s.evaluate("() => window.dispatchEvent(new Event('afterprint'))")
        r = PdfReader(str(pdf))
        ilkler = [i for i, pg in enumerate(r.pages) if "GRUBU" in (pg.extract_text() or "").replace(" ", "")]
        resim = [len(r.pages[i].images) for i in ilkler]
        kontrol(len(ilkler) == G and all(n >= 2 for n in resim), f"PDF: her grubun ilk sayfasında iki logo ({resim})")
        subprocess.run(["pdftoppm", "-r", "80", "-png", "-f", "1", "-l", "1", "-singlefile", str(pdf), str(CIKTI / "04-pdf-sayfa1")], check=True)

        # Word
        with s.expect_download(timeout=30000) as dl:
            s.locator(".sh-ust-dugme:has-text('Word')").click()
        docx = CIKTI / "sinav.docx"
        dl.value.save_as(str(docx))
        z = zipfile.ZipFile(docx)
        medya = [n for n in z.namelist() if n.startswith("word/media/")]
        belge = z.read("word/document.xml").decode("utf8")
        kontrol(len(medya) >= 2, f"Word dosyasında logo görselleri ({len(medya)})")
        kontrol(belge.count("<pic:pic") >= 2 * G or belge.count("a:blip") >= 2 * G, f"Word: her grubun başlığında iki logo ({belge.count('a:blip')} görsel)")
        try:
            subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(CIKTI), str(docx)], check=True, capture_output=True, timeout=120)
            subprocess.run(["pdftoppm", "-r", "80", "-png", "-f", "1", "-l", "1", "-singlefile", str(CIKTI / "sinav.pdf"), str(CIKTI / "05-word-sayfa1")], check=True)
        except Exception as e:
            print("   (Word -> PDF görüntüsü alınamadı:", e, ")")

        # yeni sınav: logolar kendiliğinden gelir
        s.goto(URL + "#/sinav")
        s.wait_for_selector(".sh-liste-sayfa")
        s.click("button:has-text('Yeni sınav')")
        s.click(".pencere .sh-sablon:has-text('Test')")
        s.fill(".pencere input[type=number]", "3")
        s.click("button:has-text('Sınavı oluştur')")
        s.wait_for_selector(".sh-oge-coktan")
        s.wait_for_function("() => document.querySelectorAll('.sh-logo-yeri.dolu img').length === 2", timeout=10000)
        kontrol(True, "yeni sınavda logolar hazır geldi")
        yeni_sid = s.url.split("#/sinav/")[1]
        s.locator(".sh-logo-yeri").nth(1).locator("button[aria-label$='kaldır']").click(force=True)
        s.wait_for_function("() => document.querySelectorAll('.sh-logo-yeri.dolu').length === 1", timeout=5000)
        profil = s.evaluate("() => JSON.parse(localStorage.getItem('optik-okuyucu.ogretmen') || '{}')")
        kontrol(bool(profil.get("logoSol")) and not profil.get("logoSag"), "logo kaldırılınca sonraki sınavlara da gelmez")

        print("2) Optikte soru iptali")
        s.goto(URL + f"#/sinav/{sid}/yazdir")
        s.wait_for_selector(".sh-optik-kart button:has-text('Optik okuyucuya aktar')")
        s.click(".sh-optik-kart button:has-text('Optik okuyucuya aktar')")
        s.wait_for_selector(".sh-aktar-onay:not([disabled])")
        s.click(".sh-aktar-onay")
        s.wait_for_function("() => location.hash === '#/optik'")
        durum = s.evaluate("() => JSON.parse(localStorage.getItem('optik-okuyucu.sinav.v1'))")
        anahtar = durum["anahtarlar"]
        kimlik = durum["soruKimlikleri"]
        kontrol(sorted(anahtar) == ["A", "B"] and all(len(kimlik[k]) == N for k in kimlik), "aktarıldı: A, B anahtarları ve soru eşleşmesi")

        rng = np.random.default_rng(5)
        ogrenciler = []
        adlar = [("ALİ", "KAYA"), ("AYŞE", "DEMİR"), ("CAN", "ÖZ"), ("ECE", "YILMAZ"), ("EFE", "ŞAHİN"), ("SU", "ÇELİK"), ("NUR", "AK")]
        for i, (a_, s_) in enumerate(adlar):
            kit = "AB"[i % 2]
            cev = []
            for q in range(N):
                if i == 0:
                    cev.append({"t": "c", "k": anahtar[kit][q]})          # hepsi doğru
                else:
                    x = rng.random()
                    cev.append({"t": "b"} if x < 0.15 else {"t": "x", "ks": [0, 1]} if x < 0.2 else {"t": "c", "k": int(rng.integers(0, K))})
            ogrenciler.append({"id": f"o{i}", "ad": a_, "soyad": s_, "no": str(100 + i), "kitapcik": kit, "cevaplar": cev, "notlar": [], "zaman": 0})
        durum["ogrenciler"] = ogrenciler
        durum["ekran"] = "sonuc"
        s.evaluate("d => localStorage.setItem('optik-okuyucu.sinav.v1', JSON.stringify(d))", durum)
        s.reload()
        s.wait_for_selector(".iptal-kart")

        def ui_puanlar():
            return [float(t.replace(",", ".")) for t in s.locator(".sonuc-listesi .puan").all_inner_texts()]

        def beklenen(iptaller):
            out = []
            for o in ogrenciler:
                ip = {}
                for tur, sorular in iptaller:
                    ip[sorular[o["kitapcik"]]] = tur
                out.append(puan_hesapla(o, anahtar[o["kitapcik"]], N, ip))
            return out

        kontrol(ui_puanlar() == beklenen([]), f"iptalsiz puanlar doğru {ui_puanlar()[:4]}…")
        s.screenshot(path=str(CIKTI / "06-sonuc-iptalsiz.png"), full_page=True)

        def iptal_et(kit, no, tur):
            s.click(".iptal-kart button:has-text('Soru iptal et')")
            s.select_option(".iptal-form select", kit)
            s.fill(".iptal-form input[type=number]", str(no))
            diger = "B" if kit == "A" else "A"
            q = kimlik[kit].index(kimlik[kit][no - 1])
            karsi = kimlik[diger].index(kimlik[kit][no - 1]) + 1
            yazi = s.locator(".iptal-form .kucuk-not").first.inner_text()
            kontrol(f"{diger} kitapçığında {karsi}." in yazi, f"{kit}{no} -> {diger}{karsi} otomatik bulundu ({yazi.strip()})")
            s.click(f".iptal-turler button:has-text('{'Herkese doğru say' if tur == 'dogru' else 'Soruyu çıkar'}')")
            s.screenshot(path=str(CIKTI / f"07-iptal-form-{tur}.png"), full_page=False)
            s.click(".iptal-form button:has-text('İptal et')")
            s.wait_for_selector(".iptal-form", state="detached")
            return {kit: no - 1, diger: karsi - 1}

        i1 = iptal_et("A", 3, "dogru")
        kontrol(ui_puanlar() == beklenen([("dogru", i1)]), f"'Herkese doğru say' sonrası puanlar birebir {ui_puanlar()[:4]}…")
        i2 = iptal_et("B", 7, "cikar")
        bek = beklenen([("dogru", i1), ("cikar", i2)])
        kontrol(ui_puanlar() == bek, f"'Soruyu çıkar' sonrası puanlar birebir {ui_puanlar()[:4]}…")
        kontrol(ui_puanlar()[0] == 100.0, "her soruyu doğru yapan öğrenci yine 100")
        kontrol(s.locator(".iptal-listesi li").count() == 2, "iptal listesi: 2 soru")
        # aynı soru ikinci kez iptal edilemez
        s.click(".iptal-kart button:has-text('Soru iptal et')")
        s.select_option(".iptal-form select", "A")
        s.fill(".iptal-form input[type=number]", "3")
        s.click(".iptal-form button:has-text('İptal et')")
        kontrol("zaten iptal" in s.locator(".iptal-form .hata-kutu").inner_text(), "aynı soru ikinci kez iptal edilemez")
        s.click(".iptal-form button:has-text('Vazgeç')")
        s.screenshot(path=str(CIKTI / "08-sonuc-iptalli.png"), full_page=True)

        # Excel
        with s.expect_download(timeout=30000) as dl:
            s.click("button:has-text(\"Excel'i indir\")")
        xl = CIKTI / "sonuc.xlsx"
        dl.value.save_as(str(xl))
        import openpyxl
        wb = openpyxl.load_workbook(xl)
        ws = wb["Sonuçlar"]
        bas = [c.value for c in ws[1]]
        puan_i = bas.index("Puan")
        ex = [round(float(r[puan_i]), 2) for r in ws.iter_rows(min_row=2, values_only=True)]
        kontrol(ex == bek, f"Excel puanları birebir {ex[:4]}…")
        wk = wb["Cevap Anahtarı"]
        mavi = {(row[0].value, j) for row in wk.iter_rows(min_row=2) for j, c in enumerate(row[1:]) if c.fill and c.fill.fgColor and c.fill.fgColor.rgb == "FFDDEBF7"}
        kontrol({("A", i1["A"]), ("B", i1["B"]), ("A", i2["A"]), ("B", i2["B"])} <= mavi, "Excel cevap anahtarında iptal edilen sorular işaretli")
        bilgi = {r[0]: r[1] for r in wb["Bilgi"].iter_rows(values_only=True)}
        kontrol("Herkese doğru" in str(bilgi.get("İptal edilen sorular")) and "çıkarıldı" in str(bilgi.get("İptal edilen sorular")), f"Excel bilgi: {bilgi.get('İptal edilen sorular')}")
        wa = wb["Soru Analizi"]
        ip_sutun = [c.value for c in wa[1]].index("İptal")
        iptalli = [(r[0], r[1]) for r in wa.iter_rows(min_row=2, values_only=True) if r[ip_sutun]]
        kontrol(sorted(iptalli) == sorted([("A", i1["A"] + 1), ("B", i1["B"] + 1), ("A", i2["A"] + 1), ("B", i2["B"] + 1)]), f"soru analizinde iptal sütunu {iptalli}")

        # geri al
        s.locator(".iptal-listesi li").nth(1).locator("button:has-text('Geri al')").click()
        s.locator(".iptal-listesi li").nth(0).locator("button:has-text('Geri al')").click()
        kontrol(ui_puanlar() == beklenen([]), "iptaller geri alınınca eski puanlar")

        # telefon ve karanlık görünüm
        iptal_et("A", 5, "cikar")
        s.set_viewport_size({"width": 390, "height": 844})
        s.wait_for_timeout(300)
        s.locator(".iptal-kart").screenshot(path=str(CIKTI / "10-telefon-iptal.png"))
        s.emulate_media(color_scheme="dark")
        s.wait_for_timeout(300)
        s.locator(".iptal-kart").screenshot(path=str(CIKTI / "11-telefon-karanlik.png"))
        s.set_viewport_size({"width": 1440, "height": 900})
        s.emulate_media(color_scheme="light")

        tarayici.close()

    kontrol(not konsol, "tarayıcı konsolunda hata yok" + ("" if not konsol else ": " + " | ".join(konsol[:4])))
    print(f"Çıktılar: {CIKTI}")
    print("SONUÇ:", "BAŞARILI" if not hatalar else f"{len(hatalar)} HATA")
    return 1 if hatalar else 0


if __name__ == "__main__":
    sys.exit(main())
