# Uçtan uca tarayıcı testi (canlı kamera akışı, fotoğraf yükleme olmadan):
#   ayarlar -> A ve B anahtarı kamerayla -> öğrenciler kamerayla (zor kâğıtta kontrol penceresi) -> tekrar sayılmama
#   -> sayfa yenileme -> Excel içeriği   (optik ekranı: <adres>/#/optik)
# Hazırlık:
#   python3 arac/e2e_veri.py /tmp/e2e
#   python3 arac/e2e_video.py /tmp/e2e            (v_anahtarA / v_anahtarB / v_ogrenci videoları + video_ogrenciler.json)
#   npm run build && npx vite preview   ->   python3 arac/e2e_kamera_test.py /tmp/e2e http://localhost:4173/
import json, re, shutil, sys, tempfile, time
from pathlib import Path
from playwright.sync_api import sync_playwright
import openpyxl

KOK = Path(sys.argv[1])
URL = sys.argv[2] if len(sys.argv) > 2 else "http://localhost:4173/"
OPTIK = URL.split("#")[0] + "#/optik"   # ana sayfada modül seçilir; optik ekranı bu adreste
CIKTI = Path(sys.argv[3]) if len(sys.argv) > 3 else KOK
B = json.loads((KOK / "beklenen.json").read_text())
N = B["N"]
SIK = "ABCDE"
VIDEO_OGR = json.loads((KOK / "video_ogrenciler.json").read_text())


def tr_baslik(s):
    kucuk = s.replace("I", "ı").replace("İ", "i").lower()
    out = []
    for k in kucuk.split():
        ilk = k[0]
        ilk = "İ" if ilk == "i" else "I" if ilk == "ı" else ilk.upper()
        out.append(ilk + k[1:])
    return " ".join(out)


hatalar = []
def kontrol(kosul, mesaj):
    print(("  ✔ " if kosul else "  ✘ ") + mesaj)
    if not kosul:
        hatalar.append(mesaj)


profil = Path(tempfile.mkdtemp(prefix="optik-e2e-"))
konsol = []


def ac(p, video):
    ctx = p.chromium.launch_persistent_context(
        str(profil), viewport={"width": 430, "height": 920}, accept_downloads=True, permissions=["camera"],
        args=["--use-fake-ui-for-media-stream", "--use-fake-device-for-media-stream", f"--use-file-for-fake-video-capture={video}"])
    # çalan sesleri kaydet (süre ve tepe seviyesi): okundu sesi = 275 ms "bi-bip", uyarı = 430 ms
    ctx.add_init_script("""(() => {
      window.__sesler = []
      const bas = AudioBufferSourceNode.prototype.start
      AudioBufferSourceNode.prototype.start = function (...a) {
        try { const d = this.buffer.getChannelData(0); let t = 0; for (let i = 0; i < d.length; i += 7) t = Math.max(t, Math.abs(d[i]))
          window.__sesler.push({ ms: Math.round(this.buffer.duration * 1000), tepe: +t.toFixed(2), zaman: Date.now() }) } catch (e) {}
        return bas.apply(this, a)
      }
    })()""")
    s = ctx.pages[0] if ctx.pages else ctx.new_page()
    s.on("console", lambda m: konsol.append(f"{m.type}: {m.text}") if m.type == "error" and "fonts.g" not in m.text else None)
    s.on("pageerror", lambda e: konsol.append(f"pageerror: {e}"))
    s.on("dialog", lambda d: d.accept())
    return ctx, s


def okuyucu_hazir(s):
    t = time.time()
    s.wait_for_function("() => !document.querySelector('.yukleme')", timeout=180000)
    return time.time() - t


with sync_playwright() as p:
    # ---------------------------------------------------------------- 1) ayarlar + A anahtarı
    print("1) Ayarlar ve A anahtarı (kamera)")
    ctx, s = ac(p, KOK / "v_anahtarA.mjpeg")
    t0 = time.time()
    s.goto(OPTIK)
    s.fill('input[placeholder^="Örn."]', "E2E Kamera Sınavı")
    s.fill('label:has-text("Soru sayısı") input', str(N))
    kontrol("100" in s.locator("text=Toplam:").inner_text(), "otomatik puan toplamı 100")
    s.click("text=Devam: Cevap anahtarı")
    s.wait_for_selector('button:has-text("Anahtarı kamerayla okut"):not([disabled])', timeout=180000)
    print(f"  okuyucu {time.time() - t0:.1f} sn'de hazır, mod={s.evaluate('window.__optikMod')}")

    def anahtar_oku(kit):
        s.click('button:has-text("Anahtarı kamerayla okut")')
        t = time.time()
        s.wait_for_selector("text=Anahtarı kaydet", timeout=90000)
        print(f"  {kit} anahtarı {time.time() - t:.1f} sn'de okundu")
        secili = s.locator(".anahtar-satir").evaluate_all(
            "els => els.map(e => { const s = e.querySelector('.sik.secili'); return s ? s.textContent : null })")
        beklenen = [SIK[B["anahtarlar"][kit][str(q)]] for q in range(N)]
        kontrol(secili == beklenen, f"{kit} anahtarının {N} sorusu doğru okundu")
        kontrol(s.locator(".pencere .sec.secili").first.inner_text() == kit, f"{kit} kitapçığı otomatik seçildi")
        kontrol(not s.is_visible("text=Boş sorular"), "anahtarda eksik soru yok")
        s.screenshot(path=str(CIKTI / f"e2e-anahtar-{kit}.png"))
        s.click("text=Anahtarı kaydet")
        s.wait_for_selector(f"text=Kitapçık {kit}")

    anahtar_oku("A")
    ctx.close()

    # ---------------------------------------------------------------- 2) B anahtarı
    print("2) B anahtarı (kamera, tarayıcı yeniden açıldı: kayıt korunmalı)")
    ctx, s = ac(p, KOK / "v_anahtarB.mjpeg")
    s.goto(OPTIK)
    kontrol(s.is_visible("text=Kitapçık A"), "A anahtarı yeniden açılışta duruyor")
    okuyucu_hazir(s)
    anahtar_oku("B")
    s.click("text=Devam: Öğrenci kâğıtlarını okut")
    ctx.close()

    # ---------------------------------------------------------------- 3) öğrenciler
    print("3) Öğrenci kâğıtları (canlı kamera)")
    ctx, s = ac(p, KOK / "v_ogrenci.mjpeg")
    s.goto(OPTIK)
    okuyucu_hazir(s)
    if s.is_visible("text=Kapat"):
        s.click("text=Kapat")   # "nasıl okutulur" kutusu
    s.click('button:has-text("Kamerayla okut")')
    s.wait_for_selector("video")
    cozunurluk = s.locator(".kamera-coz").inner_text(timeout=20000) if s.locator(".kamera-coz").count() else "?"
    print("  kamera çözünürlüğü:", cozunurluk)
    beklenenler = [o for o in B["ogrenciler"] if o["dosya"] in VIDEO_OGR]
    bildirimler, zamanlar, mesajlar = [], [], set()
    pencere = 0
    t0 = time.time()
    ekran_tamam = ekran_kontrol = False
    son_sayi = 0
    while time.time() - t0 < 110:
        try:
            if s.is_visible("text=Onayla ve kaydet"):
                no_input = s.locator('.form-izgara label:has-text("Numara") input')
                ad_in = s.locator('.form-izgara label:has-text("Ad") input').first
                okunan_no, okunan_ad = no_input.input_value(), ad_in.input_value()
                ogr = next((o for o in beklenenler if o["no"] == okunan_no), None) or next((o for o in beklenenler if o["ad"] == okunan_ad), None)
                if ogr is None:
                    desen = re.compile("^" + re.escape(okunan_no).replace("\\?", ".") + "$")
                    ogr = next((o for o in beklenenler if desen.match(o["no"])), None)
                if not ekran_kontrol:
                    time.sleep(0.6); s.screenshot(path=str(CIKTI / "e2e-kontrol-penceresi.png")); ekran_kontrol = True
                zaten = s.locator(".sonuc-listesi li .isim small").evaluate_all("els => els.map(e => (e.textContent.match(/No (\\d+)/) || [])[1])")
                if ogr is None or ogr["no"] in zaten:
                    s.click("text=Bu kâğıdı atla"); continue
                pencere += 1
                print(f"  · kontrol penceresi: {ogr['dosya']} -> {s.locator('.sorun-listesi').inner_text().strip().splitlines()[:2]}")
                ad_in.fill(ogr["ad"]); s.locator('.form-izgara label:has-text("Soyad") input').fill(ogr["soyad"]); no_input.fill(ogr["no"])
                for blok in s.locator(".soru-kontrol").all():
                    q = int(blok.locator(".soru-baslik").inner_text().split(".")[0]) - 1
                    c = ogr["cevaplar"][str(q)]
                    if c["tur"] == "cift": blok.locator("button:has-text('Çift')").click()
                    elif c["tur"] == "bos": blok.locator("button:has-text('Boş')").click()
                    else: blok.locator(f"button.sec:text-is('{SIK[c['k']]}')").click()
                s.click("text=Onayla ve kaydet")
                continue
            mesajlar.add(s.locator(".kamera-mesaj").inner_text(timeout=800))
            if s.locator(".bildirim").count():
                b = s.locator(".bildirim").inner_text(timeout=500)
                if not bildirimler or bildirimler[-1] != b:
                    bildirimler.append(b); zamanlar.append(round(time.time() - t0, 1))
                    if "✅" in b and not ekran_tamam:
                        s.screenshot(path=str(CIKTI / "e2e-kamera-okundu.png")); ekran_tamam = True
            sayi = int(s.locator(".sayac").inner_text(timeout=500).split()[0])
            if sayi != son_sayi:
                son_sayi = sayi
                if sayi == len(beklenenler):
                    print(f"  tüm kâğıtlar {time.time() - t0:.1f} sn'de okundu; tekrar sayılmadığı kontrol ediliyor…")
                    t_bitti = time.time()
            if son_sayi == len(beklenenler) and time.time() - t_bitti > 32:
                break
        except Exception:
            pass
        time.sleep(0.2)
    print("  bildirimler:", list(zip(zamanlar, bildirimler)))
    print("  görülen kamera mesajları:", sorted(mesajlar))
    sesler = s.evaluate("window.__sesler")
    okundu_sesi = [x for x in sesler if x["ms"] == 275]
    uyari_sesi = [x for x in sesler if x["ms"] == 430]
    print(f"  sesler: {len(okundu_sesi)} okundu, {len(uyari_sesi)} uyarı; tepe {sorted({x['tepe'] for x in sesler})}")
    kaydedilen = int(s.locator(".sayac").inner_text().split()[0])
    kontrol(kaydedilen == len(okundu_sesi), f"her kaydedilen kâğıtta bir kez yüksek 'okundu' sesi ({len(okundu_sesi)} ses / {kaydedilen} kâğıt)")
    kontrol(all(x["tepe"] >= 0.9 for x in sesler), "sesler tam seviyeye yakın (yüksek)")
    kontrol(len(uyari_sesi) >= pencere + (1 if any("zaten okunmuş" in b for b in bildirimler) else 0), f"kontrol penceresi / aynı kâğıt için ayrı uyarı sesi ({len(uyari_sesi)})")
    # ses düğmesi: kapat -> hatırlanır; aç -> örnek ses
    s.click(".kamera-ses")
    kontrol(s.get_attribute(".kamera-ses", "aria-pressed") == "false" and s.evaluate("localStorage.getItem('optik-okuyucu.ses')") == "kapali", "ses düğmesi: kapatıldı ve hatırlandı")
    n0 = s.evaluate("window.__sesler.length")
    s.click(".kamera-ses")
    kontrol(s.get_attribute(".kamera-ses", "aria-pressed") == "true" and s.evaluate("window.__sesler.length") == n0 + 1, "ses düğmesi: açılınca örnek ses çaldı")
    s.locator(".kamera-ekran-ust").screenshot(path=str(CIKTI / "e2e-kamera-ses-dugmesi.png"))
    s.click(".kamera-ekran-kapat")

    def liste_oku():
        return s.locator(".sonuc-listesi li").evaluate_all(
            "els => els.map(e => [e.querySelector('.isim').childNodes[0].textContent.trim(), e.querySelector('.puan').textContent.trim(), e.querySelector('.isim small').textContent])")

    liste = liste_oku()
    bek = {(tr_baslik(o["ad"] + " " + o["soyad"]), o["no"], f"{o['puan']:.2f}".replace(".", ",")) for o in beklenenler}
    gorulen = {(a, (re.search(r"No (\d+)", k) or [None, None])[1], p_) for a, p_, k in liste}
    kontrol(len(liste) == len(beklenenler), f"{len(beklenenler)} öğrenci okundu, döngüde tekrar sayılmadı (listede {len(liste)})")
    kontrol(gorulen == bek, "isimler, numaralar ve puanlar birebir doğru")
    if gorulen != bek:
        print("   eksik:", bek - gorulen); print("   fazla:", gorulen - bek)
    kontrol(pencere >= 1, f"zor kâğıtta kontrol penceresi açıldı ({pencere})")
    kontrol(any("✅" in b for b in bildirimler), "okununca '✅ Ad Soyad — puan' bildirimi çıktı")
    kontrol(any("zaten okunmuş" in b for b in bildirimler), "video başa dönünce 'zaten okunmuş' uyarısı çıktı")

    print("4) Sayfa yenileme")
    s.reload()
    s.wait_for_selector(".sonuc-listesi li")
    kontrol(len(liste_oku()) == len(beklenenler), "yenilemeden sonra liste duruyor")
    s.locator(".sonuc-listesi li").first.click()
    time.sleep(0.5)
    s.screenshot(path=str(CIKTI / "e2e-ogrenci-detay.png"))
    s.locator(".pencere button:has-text('Kapat')").last.click()

    print("5) Excel")
    s.click('button:has-text("Bitti")')
    with s.expect_download() as indir:
        s.click("text=Excel'i indir")
    yol = CIKTI / "e2e-indirilen.xlsx"
    indir.value.save_as(str(yol))
    wb = openpyxl.load_workbook(yol)
    kontrol(wb.sheetnames == ["Sonuçlar", "Sıralama", "Cevaplar", "Soru Analizi", "Cevap Anahtarı", "Bilgi"], f"Excel sayfaları: {wb.sheetnames}")
    satirlar = list(wb["Sonuçlar"].iter_rows(min_row=2, values_only=True))
    excel = {(str(r[1]), f"{r[2]} {r[3]}", r[4], r[5], r[6], r[7], round(r[9], 2)) for r in satirlar}
    bekx = {(o["no"], tr_baslik(o["ad"] + " " + o["soyad"]), o["kitapcik"], o["d"], o["y"], o["b"], o["puan"]) for o in beklenenler}
    kontrol(excel == bekx, "Excel satırları (no, ad soyad, kitapçık, D/Y/B, puan) birebir doğru")
    anahtar_satir = {r[0]: "".join(r[1:]) for r in wb["Cevap Anahtarı"].iter_rows(min_row=2, values_only=True)}
    kontrol(anahtar_satir == {k: "".join(SIK[v[str(q)]] for q in range(N)) for k, v in B["anahtarlar"].items()}, "Excel'deki cevap anahtarları doğru")

    s.screenshot(path=str(CIKTI / "e2e-sonuc.png"), full_page=True)
    kontrol(not konsol, f"tarayıcı konsolunda hata yok {konsol[:3]}")
    ctx.close()

shutil.rmtree(profil, ignore_errors=True)
print("\nSONUÇ:", "TÜM KONTROLLER GEÇTİ" if not hatalar else f"{len(hatalar)} HATA: {hatalar}")
sys.exit(1 if hatalar else 0)
